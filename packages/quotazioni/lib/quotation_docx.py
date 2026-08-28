# quotation_docx.py

from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
from docx import Document

import re
import base64
import tempfile

from html import unescape

from docxtpl import RichText
from htmldocx import HtmlToDocx

from copy import deepcopy


class QuotationDocxBuilder(object):

    def __init__(self, db, site):

        self.db = db
        self.site = site

        self.tbl_quot = self.db.table(
            'quotazioni.quotazione_versione'
        )

    # ---------------------------------------------------------
    # Converte una relazione Genropy in lista per docxtpl
    # ---------------------------------------------------------

    def bag_rows(self, bag, relation, mapping):

        result = []

        rows = bag.getItem(relation)

        if not rows:
            return result

        for node in rows:

            row = node.value

            item = {}

            for key, field in mapping.items():

                if row:
                    item[key] = row.getItem(field) or ''
                else:
                    item[key] = ''

            result.append(item)

        return result

    # ---------------------------------------------------------
    # HTML -> testo
    # ---------------------------------------------------------

    def html_to_text(self, value):

        if not value:
            return ''

        value = unescape(value)

        value = value.replace(
            '\t',
            ''
        )

        value = re.sub(
            r'<\s*br\s*/?>',
            '\n',
            value,
            flags=re.I
        )

        value = re.sub(
            r'</\s*div\s*>',
            '\n',
            value,
            flags=re.I
        )

        value = re.sub(
            r'<\s*div[^>]*>',
            '',
            value,
            flags=re.I
        )

        value = re.sub(
            r'<[^>]+>',
            '',
            value
        )

        value = value.replace(
            '\xa0',
            ' '
        )

        value = '\n'.join(
            line.strip()
            for line in value.split('\n')
        )

        value = re.sub(
            r'\n\s*\n\s*\n+',
            '\n\n',
            value
        )

        return value.strip()

    # ---------------------------------------------------------
    # Pulizia testo
    # ---------------------------------------------------------

    def clean_text(self, value):

        if not value:
            return ''

        value = value.replace(
            '\r\n',
            '\n'
        )

        value = re.sub(
            r'\n\s*\n\s*\n+',
            '\n\n',
            value
        )

        return value.strip()

    # ---------------------------------------------------------
    # Rows -> testo
    # ---------------------------------------------------------

    def rows_to_text(
        self,
        rows,
        bullet=False
    ):

        if not rows:
            return ''

        result = []

        for row in rows:

            description = row.get(
                'description',
                ''
            )

            if not description:
                continue

            if bullet:

                result.append(
                    f'• {description}'
                )

            else:

                result.append(
                    description
                )

        return '\n'.join(result)

    # ---------------------------------------------------------
    # Rows -> RichText
    # ---------------------------------------------------------

    def rows_to_richtext(
        self,
        rows,
        bullet=False
    ):

        rt = RichText()

        if not rows:
            return rt

        first = True

        for row in rows:

            text = row.get(
                'description',
                ''
            )

            if not text:
                continue

            if not first:

                rt.add(
                    '\n'
                )

            if bullet:

                rt.add(
                    '• '
                )

            rt.add(
                text
            )

            first = False

        return rt

    # ---------------------------------------------------------
    # Cerca il placeholder BODY_PLACEHOLDER
    #
    # Cerca sia nei paragrafi normali sia nelle celle
    # delle eventuali tabelle del template.
    # ---------------------------------------------------------

    def find_body_placeholder(
        self,
        container
    ):

        # -------------------------------------------------
        # Paragrafi
        # -------------------------------------------------

        for paragraph in container.paragraphs:

            if 'BODY_PLACEHOLDER' in paragraph.text:

                return paragraph

        # -------------------------------------------------
        # Tabelle
        # -------------------------------------------------

        for table in container.tables:

            for row in table.rows:

                for cell in row.cells:

                    result = self.find_body_placeholder(
                        cell
                    )

                    if result:

                        return result

        return None

    # ---------------------------------------------------------
    # Inserisce HTML nel punto del placeholder
    # ---------------------------------------------------------

    def insert_html_at_placeholder(
        self,
        placeholder_paragraph,
        html
    ):

        if not placeholder_paragraph:
            return False

        if not html:
            return False

        # -------------------------------------------------
        # Documento Word temporaneo
        # -------------------------------------------------

        temp_doc = Document()

        parser = HtmlToDocx()

        parser.add_html_to_document(
            html,
            temp_doc
        )

        # -------------------------------------------------
        # Recuperiamo gli elementi XML generati
        # -------------------------------------------------

        temp_body = temp_doc.element.body

        elements = []

        for element in temp_body:

            # Non copiare sectPr
            if element.tag.endswith(
                'sectPr'
            ):
                continue

            elements.append(
                deepcopy(element)
            )

        if not elements:
            return False

        # -------------------------------------------------
        # Paragrafo placeholder
        # -------------------------------------------------

        placeholder_element = (
            placeholder_paragraph._p
        )

        parent = (
            placeholder_element.getparent()
        )

        if parent is None:
            return False

        # -------------------------------------------------
        # Posizione del placeholder
        # -------------------------------------------------

        index = parent.index(
            placeholder_element
        )

        # -------------------------------------------------
        # Inseriamo gli elementi HTML
        # prima del placeholder
        # -------------------------------------------------

        for offset, element in enumerate(
            elements
        ):

            parent.insert(
                index + offset,
                element
            )

        # -------------------------------------------------
        # Eliminiamo il placeholder
        # -------------------------------------------------

        parent.remove(
            placeholder_element
        )

        return True

    # ---------------------------------------------------------
    # Costruzione context
    # ---------------------------------------------------------

    def build_context(
        self,
        rec
    ):

        ctx = {}

        # -------------------------------------------------
        # Campi principali
        # -------------------------------------------------

        main_fields = [
            'data',
            'quot_n',
            'oggetto',
            'body_in',
            'note',
            'servincl_int',
            'servexcl_int',
            'spec_cond_int',
            'surcharges_int',
            'workingtimes_int',
            'agent_sign',
            'agent_name'
        ]

        for field in main_fields:

            value = rec.getItem(
                field
            ) or ''

            # body_in deve rimanere HTML
            if field == 'body_in':

                value = value or ''

            ctx[field] = value

        ctx['agency_logo'] = ''
        ctx['stamp'] = ''

        # -------------------------------------------------
        # Cliente
        # -------------------------------------------------

        ctx['customer_rag_sociale'] = (
            rec['@quotazione_id'].getItem(
                '@cliente_id.rag_sociale'
            ) or ''
        )

        ctx['customer_address'] = (
            rec['@quotazione_id'].getItem(
                '@cliente_id.address'
            ) or ''
        )

        ctx['customer_cap'] = (
            rec['@quotazione_id'].getItem(
                '@cliente_id.cap'
            ) or ''
        )

        ctx['customer_city'] = (
            rec['@quotazione_id'].getItem(
                '@cliente_id.city'
            ) or ''
        )

        # -------------------------------------------------
        # Agency
        # -------------------------------------------------

        ctx['agency_name'] = (
            rec['@quotazione_id'].getItem(
                '@agency_id.agency_name'
            ) or ''
        )

        ctx['agency_address'] = (
            rec['@quotazione_id'].getItem(
                '@agency_id.address'
            ) or ''
        )

        ctx['agency_tel'] = (
            rec['@quotazione_id'].getItem(
                '@agency_id.tel'
            ) or ''
        )

        ctx['agency_fax'] = (
            rec['@quotazione_id'].getItem(
                '@agency_id.fax'
            ) or ''
        )

        ctx['agency_email'] = (
            rec['@quotazione_id'].getItem(
                '@agency_id.email'
            ) or ''
        )

        ctx['agency_web'] = (
            rec['@quotazione_id'].getItem(
                '@agency_id.web'
            ) or ''
        )

        ctx['agency_vat'] = (
            rec['@quotazione_id'].getItem(
                '@agency_id.vat'
            ) or ''
        )

        ctx['agency_cf'] = (
            rec['@quotazione_id'].getItem(
                '@agency_id.cf'
            ) or ''
        )

        # -------------------------------------------------
        # Data
        # -------------------------------------------------

        data_value = rec.getItem(
            'data'
        )

        if data_value:

            ctx['data'] = data_value.strftime(
                "%d/%m/%Y"
            )

        else:

            ctx['data'] = ''

        # -------------------------------------------------
        # Numero quotazione
        # -------------------------------------------------
        
        quot_n = (rec['@quotazione_id'].getItem(
                        'quot_n'
                    ) or ''
                )
        vers_n = (rec.getItem(
                        'version'
                    ) or ''
                )
        
        ctx['quot_n'] = ('%s_%s'
                         % (quot_n,vers_n) or ''
        )

        # -------------------------------------------------
        # Tabelle
        # -------------------------------------------------

        ctx['description_quot'] = self.bag_rows(
            rec,
            '@description_quot',
            {
                'description': 'description',
                'um': 'um',
                'amount': 'amount'
            }
        )

        ctx['quotazione_extra'] = self.bag_rows(
            rec,
            '@quotazione_extra',
            {
                'description': 'extra_description',
                'um': 'extra_um',
                'amount': 'extra_amount'
            }
        )

        ctx['note_quot'] = self.bag_rows(
            rec,
            '@note_quot',
            {
                'description': 'description',
                'um': 'um',
                'amount': 'amount'
            }
        )

        # -------------------------------------------------
        # Servizi inclusi
        # -------------------------------------------------

        rows = self.bag_rows(
            rec,
            '@servincl_quot',
            {
                'description':
                    '@servicesincl_id.description'
            }
        )

        ctx['servincl_text'] = (
            self.rows_to_text(
                rows,
                bullet=True
            )
        )

        # -------------------------------------------------
        # Servizi esclusi
        # -------------------------------------------------

        rows = self.bag_rows(
            rec,
            '@serviexcl_quot',
            {
                'description':
                    '@servicesexcl_id.description'
            }
        )

        ctx['serviexcl_text'] = (
            self.rows_to_text(
                rows,
                bullet=True
            )
        )

        # -------------------------------------------------
        # Condizioni specifiche
        # -------------------------------------------------

        rows = self.bag_rows(
            rec,
            '@specific_quot',
            {
                'description':
                    'description'
            }
        )

        ctx['specific_text'] = (
            self.rows_to_text(
                rows
            )
        )

        # -------------------------------------------------
        # Maggiorazioni
        # -------------------------------------------------

        rows = self.bag_rows(
            rec,
            '@surcharge_quot',
            {
                'description':
                    '@surcharges_id.description'
            }
        )

        ctx['surcharge_text'] = (
            self.rows_to_text(
                rows,
                bullet=True
            )
        )

        # -------------------------------------------------
        # Orari
        # -------------------------------------------------

        rows = self.bag_rows(
            rec,
            '@times_quot',
            {
                'description':
                    '@worktimes_id.description'
            }
        )

        ctx['times_text'] = (
            self.rows_to_text(
                rows,
                bullet=True
            )
        )

        # -------------------------------------------------
        # Accettazione da preferenze
        # -------------------------------------------------

        ctx['accettazione'] = self.html_to_text(
            self.db.application.getPreference('accettazione',
                pkg='quotazioni') or ''
        )

                

        # -------------------------------------------------
        # Pagamenti
        # -------------------------------------------------

        rows = self.bag_rows(
            rec,
            '@paym_quot',
            {
                'description':
                    '@paymentcond_id.description'
            }
        )

        ctx['paym_text'] = (
            self.rows_to_text(
                rows,
                bullet=True
            )
        )

        return ctx

    # ---------------------------------------------------------
    # Base64 -> immagine temporanea
    # ---------------------------------------------------------

    def base64_image(
        self,
        image_data
    ):

        if not image_data:
            return None

        if ',' in image_data:

            header, encoded = (
                image_data.split(
                    ',',
                    1
                )
            )

        else:

            encoded = image_data

        image_bytes = base64.b64decode(
            encoded
        )

        tmp = tempfile.NamedTemporaryFile(
            delete=False,
            suffix='.png'
        )

        tmp.write(
            image_bytes
        )

        tmp.close()

        return tmp.name

    # ---------------------------------------------------------
    # Generazione documento
    # ---------------------------------------------------------

    def build(
        self,
        record_id
    ):

        # -------------------------------------------------
        # Record
        # -------------------------------------------------

        rec = self.tbl_quot.record(
            record_id
        ).output(
            'bag'
        )

        context = self.build_context(
            rec
        )

        # -------------------------------------------------
        # Template
        # -------------------------------------------------

        template_node = self.site.storageNode(
            'home:stampe_template',
            'quotation_template.docx'
        )

        template_path = (
            template_node.internal_path
        )

        doc = DocxTemplate(
            template_path
        )

        # -------------------------------------------------
        # Valori iniziali immagini
        # -------------------------------------------------

        context['agency_logo'] = ''
        context['stamp'] = ''
        context['show_signature'] = False

        # -------------------------------------------------
        # Logo agenzia
        # -------------------------------------------------

        agency_bag = (
            rec[
                '@quotazione_id.@agency_id'
            ]
        )

        logo = agency_bag.getItem(
            'agency_logo'
        )

        logo_file = self.base64_image(
            logo
        )

        if logo_file:

            context['agency_logo'] = InlineImage(
                doc,
                logo_file,
                width=Mm(75)
            )

        # -------------------------------------------------
        # Firma / timbro
        # -------------------------------------------------

        firma = rec.getItem(
            'firma'
        )

        if firma:

            context['show_signature'] = True

            stamp = agency_bag.getItem(
                'agency_stamp'
            )

            stamp_file = self.base64_image(
                stamp
            )

            if stamp_file:

                context['stamp'] = InlineImage(
                    doc,
                    stamp_file,
                    width=Mm(30)
                )

        # -------------------------------------------------
        # BODY HTML
        # -------------------------------------------------

        body_in = context.get(
            'body_in',
            ''
        )

        # -------------------------------------------------
        # IMPORTANTE:
        #
        # Non lasciamo che docxtpl sostituisca body_in.
        # Il placeholder BODY_PLACEHOLDER rimane nel
        # documento dopo il render.
        # -------------------------------------------------

        context['body_in'] = ''

        # -------------------------------------------------
        # Render normale
        # -------------------------------------------------

        doc.render(
            context
        )

        # -------------------------------------------------
        # Inserimento BODY HTML
        #
        # Cerchiamo BODY_PLACEHOLDER DOPO il render.
        # -------------------------------------------------

        if body_in:

            body_placeholder = (
                self.find_body_placeholder(
                    doc.docx
                )
            )

            if body_placeholder:

                self.insert_html_at_placeholder(
                    body_placeholder,
                    body_in
                )

        # -------------------------------------------------
        # Nome file
        # -------------------------------------------------

        quot_n = (
            rec[
                '@quotazione_id'
            ].getItem(
                'quot_n'
            ) or ''
        )

        quot_n = quot_n.replace(
            '/',
            '_'
        )

        vers_n = (
            rec.getItem(
                'version'
            ) or ''
        )

        vers_n = vers_n.replace(
            '/',
            '_'
        )

        filename = (
            'quotation_%s_%s.docx'
            % (
                quot_n,
                vers_n
            )
        )

        # -------------------------------------------------
        # Salvataggio
        # -------------------------------------------------

        output = self.site.storageNode(
            'home:stampe_template',
            filename
        )

        doc.save(
            output.internal_path
        )

        return output
